"""
Project: RFID-Based Smart Attendance System
Component: Automatic Excel Logger (Host Daemon)

Description:
Interfaces with an ESP32 microcontroller over USB serial,
listens for formatted attendance records, checks for calendar-day
duplicates, and automatically appends styled entries to Microsoft Excel.

Protocol Format:
EXCEL|<UID>|<Student Name>|PRESENT

Author: Aditya Kumar Sah
"""

import os
import sys
import time
from datetime import datetime

import serial
from serial import SerialException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# =========================
# Configuration Parameters
# =========================
SERIAL_PORT = "COM3"  # Adjust to match your ESP32 COM port (e.g., COM3, COM4, /dev/ttyUSB0)
BAUD_RATE = 115200
EXCEL_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "attendance.xlsx")
SHEET_NAME = "Attendance"
HEADERS = ["Date", "Time", "Student Name", "RFID UID", "Status"]


# =========================
# Workbook & Sheet Management
# =========================

def create_workbook_if_needed():
    """Create a new formatted workbook with header row if the file does not exist."""
    if not os.path.exists(EXCEL_FILE):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_NAME
        sheet.append(HEADERS)
        format_sheet(sheet)
        workbook.save(EXCEL_FILE)
        workbook.close()
        print(f"Created workbook: {EXCEL_FILE}")


def format_sheet(sheet):
    """Apply styling, column widths, and freeze panes to the worksheet."""
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = [14, 12, 24, 22, 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:E{max(sheet.max_row, 1)}"


# =========================
# Serial Line Parsing
# =========================

def parse_attendance_line(line):
    """
    Parse an incoming serial line.
    Returns (uid, name, status) if line matches the EXCEL| format, else None.
    """
    if not line.startswith("EXCEL|"):
        return None

    parts = line.split("|")
    if len(parts) != 4:
        print(f"Ignoring malformed attendance record: {line!r}")
        return None

    _, uid, name, status = parts
    uid = uid.strip().upper()
    name = name.strip()
    status = status.strip().upper()

    if not uid or not name or status != "PRESENT":
        print(f"Ignoring invalid attendance record: {line!r}")
        return None

    return uid, name, status


# =========================
# Duplicate Prevention & Persistence
# =========================

def already_recorded(sheet, date_text, uid):
    """Check if the given UID has already been recorded for the specified date."""
    for row in sheet.iter_rows(min_row=2, values_only=True):
        row_date, _row_time, _name, row_uid, _status = row[:5]
        if str(row_date).strip() == date_text and str(row_uid).strip().upper() == uid:
            return True
    return False


def save_attendance(uid, name, status):
    """
    Append an attendance record to the workbook if not already recorded today.
    Returns True if appended, False if duplicate.
    """
    now = datetime.now()
    date_text = now.strftime("%Y-%m-%d")
    time_text = now.strftime("%H:%M:%S")

    create_workbook_if_needed()
    workbook = load_workbook(EXCEL_FILE)

    try:
        if SHEET_NAME in workbook.sheetnames:
            sheet = workbook[SHEET_NAME]
        else:
            sheet = workbook.create_sheet(SHEET_NAME)
            sheet.append(HEADERS)
            format_sheet(sheet)

        if already_recorded(sheet, date_text, uid):
            print(f"Already recorded today: {name} (UID {uid})")
            return False

        sheet.append([date_text, time_text, name, uid, status])
        format_sheet(sheet)
        workbook.save(EXCEL_FILE)
        print(f"Saved: {name} | {uid} | {date_text} {time_text} | {status}")
        return True
    finally:
        workbook.close()


# =========================
# Main Execution Loop
# =========================

def main():
    create_workbook_if_needed()
    print("=" * 44)
    print("  RFID Attendance Automatic Excel Logger   ")
    print("=" * 44)
    print(f"Excel workbook: {EXCEL_FILE}")
    print(f"Serial port:    {SERIAL_PORT} @ {BAUD_RATE} baud")
    print("Note: Close Arduino Serial Monitor before running.")
    print("Listening for scan events. Press Ctrl+C to terminate.\n")

    try:
        with serial.Serial(SERIAL_PORT, BAUD_RATE, timeout=1) as esp:
            time.sleep(2)  # Allow ESP32 serial bridge to stabilize
            esp.reset_input_buffer()
            print("Connected to ESP32. Ready to scan cards...\n")

            while True:
                raw_line = esp.readline()
                if not raw_line:
                    continue

                line = raw_line.decode("utf-8", errors="replace").strip()
                if not line:
                    continue

                # Echo diagnostic feedback to the terminal
                print(f"ESP32: {line}")

                record = parse_attendance_line(line)
                if record is None:
                    continue

                uid, name, status = record
                try:
                    save_attendance(uid, name, status)
                except PermissionError:
                    print(
                        "ERROR: Cannot save attendance.xlsx because the file is currently "
                        "open in Microsoft Excel. Please close it and rescan the card."
                    )
                except OSError as error:
                    print(f"ERROR: Filesystem error saving workbook: {error}")

    except SerialException as error:
        print(f"\nSerial connection error: {error}")
        print("Please verify SERIAL_PORT, ensure the USB cable is connected,")
        print("and confirm that the Arduino Serial Monitor is closed.")
        sys.exit(1)
    except KeyboardInterrupt:
        print("\nAttendance logger terminated by user.")


if __name__ == "__main__":
    main()
